#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
مرحله ۵ — ساخت متن غنی‌شده (Embedding_Text) برای هر کالا
------------------------------------
شرح خام کالا را با دسته‌بندی و کلمات کلیدی ترکیب می‌کند تا کیفیت
Embedding نهایی (مرحله ۶) بهتر شود.

ورودی:  outputs/categorized_goods.xlsx
خروجی:  outputs/embedding_data.xlsx
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from common.config import OUTPUT_DIR

IN_PATH = OUTPUT_DIR / "categorized_goods.xlsx"
OUT_PATH = OUTPUT_DIR / "embedding_data.xlsx"

TEMPLATE = """نام کالا: {desc}

دسته اصلی: {l1}
زیرگروه: {l2}
رده تخصصی: {l3}

کلمات مرتبط:
{keywords}"""


def main():
    if not IN_PATH.exists():
        print("❌ ابتدا مرحله ۴ (دسته‌بندی) را اجرا کنید.")
        return

    wb_in = openpyxl.load_workbook(IN_PATH, read_only=True, data_only=True)
    ws_in = wb_in.active
    rows = list(ws_in.iter_rows(min_row=2, values_only=True))

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "EmbeddingData"
    ws_out.append(["StuffCod", "Description", "Embedding_Text"])

    for sid, desc, typ, l1, l2, l3, keywords in rows:
        kw_lines = "\n".join(k.strip() for k in (keywords or "").split(",") if k.strip())
        text = TEMPLATE.format(desc=desc or "", l1=l1 or "", l2=l2 or "", l3=l3 or "", keywords=kw_lines)
        ws_out.append([sid, desc, text])

    wb_out.save(OUT_PATH)
    print(f"✅ {len(rows):,} متن غنی‌شده ساخته شد → {OUT_PATH}")


if __name__ == "__main__":
    main()
