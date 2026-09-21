# -*- coding: utf-8 -*-
"""
مشترکات جستجوی گروهی از فایل اکسل/CSV
------------------------------------
هم روت‌های آپلود (server/app.py، برای اعتبارسنجی سریع هنگام آپلود) و هم
پردازش پس‌زمینه (pipeline/bulk_lookup.py) از همین ماژول برای خواندن/نوشتن
فایل و تعریف ستون‌ها استفاده می‌کنند تا هر دو دقیقاً یک قرارداد را ببینند.
"""
import csv
from pathlib import Path

import openpyxl

MAX_ROWS = 1001
RETENTION_DAYS = 3  # فایل ورودی/نتیجه‌ی هر job بعد از این مدت از دیسک پاک می‌شود

COL_ROW = "ردیف"
COL_DESC = "شرح کالا"
COL_BRAND = "برند"
COL_TYPE = "نوع"
COL_CATEGORY_HINT = "دسته کالا"
COL_NOTES = "توضیحات بیشتر"
COL_INTERNAL_CODE = "کد داخلی شما"

INPUT_COLUMNS = [
    COL_ROW, COL_DESC, COL_BRAND, COL_TYPE,
    COL_CATEGORY_HINT, COL_NOTES, COL_INTERNAL_CODE,
]
REQUIRED_COLUMN = COL_DESC

TYPE_IMPORTED = "وارداتی"
TYPE_DOMESTIC = "تولید داخل"
VALID_TYPES = {TYPE_IMPORTED, TYPE_DOMESTIC}

COL_OUT_ID = "شناسه کالا"
COL_OUT_OFFICIAL_DESC = "شرح رسمی شناسه"
COL_OUT_CATEGORY = "دسته"
COL_OUT_OTHER_OPTIONS = "گزینه‌های دیگر"
COL_OUT_STATUS = "وضعیت"
COL_OUT_USER_CHOICE = "انتخاب شما"

# «درصد اطمینان» عمداً حذف شد: مقدارش امتیاز خام RRF/BM25 بود، نه یک اطمینان
# کالیبره‌شده، و بین پرس‌وجوهای مختلف قابل مقایسه نبود — یک ستون با این عنوان
# کاربر را گمراه می‌کرد. «وضعیت» (یافت شد/نیاز به بررسی/یافت نشد) سیگنال واقعی
# را می‌دهد.
OUTPUT_COLUMNS = [
    COL_OUT_ID, COL_OUT_OFFICIAL_DESC, COL_OUT_CATEGORY,
    COL_OUT_OTHER_OPTIONS, COL_OUT_STATUS, COL_OUT_USER_CHOICE,
]

STATUS_FOUND = "یافت شد"
STATUS_REVIEW = "نیاز به بررسی"
STATUS_NOT_FOUND = "یافت نشد"


class ValidationError(Exception):
    """خطای اعتبارسنجی فایل، با پیام آماده برای نمایش به کاربر."""


def _read_xlsx_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(c).strip() if c is not None else "" for c in header_row]
    rows = []
    for raw in rows_iter:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        values = list(raw) + [None] * (len(headers) - len(raw))
        rec = {h: ("" if v is None else str(v).strip()) for h, v in zip(headers, values) if h}
        rows.append(rec)
    return headers, rows


def _read_csv_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        rows = []
        for raw in reader:
            if not any((v or "").strip() for v in raw.values()):
                continue
            rec = {(h or "").strip(): (raw.get(h) or "").strip() for h in raw}
            rows.append(rec)
        return headers, rows


def read_rows(path):
    """خواندن هدر و ردیف‌های فایل ورودی (.xlsx یا .csv). هر ردیف یک dict
    هدر->مقدار است؛ ردیف‌های کاملاً خالی رد می‌شوند."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return _read_xlsx_rows(path)
    if ext == ".csv":
        return _read_csv_rows(path)
    raise ValidationError(f"پسوند فایل پشتیبانی نمی‌شود: «{ext}». فقط xlsx. و csv. مجاز است. "
                           "(نکته: xls. قدیمی پشتیبانی نمی‌شود؛ در اکسل با «Save As» به xlsx. تبدیل کنید.)")


def validate_rows(headers, rows):
    """اعتبارسنجی اولیه‌ی فایل؛ در صورت خطا ValidationError با پیام فارسی روشن می‌اندازد."""
    if REQUIRED_COLUMN not in headers:
        raise ValidationError(
            f"ستون الزامی «{REQUIRED_COLUMN}» در فایل پیدا نشد. "
            "از «دانلود قالب اکسل» برای ساختار درست استفاده کنید."
        )

    if len(rows) > MAX_ROWS:
        raise ValidationError(
            f"فایل {len(rows):,} ردیف دارد؛ حداکثر {MAX_ROWS} ردیف در هر فایل مجاز است."
        )

    if not rows:
        raise ValidationError("فایل هیچ ردیف داده‌ای ندارد.")

    bad_rows = []
    for i, rec in enumerate(rows, start=1):
        if not (rec.get(REQUIRED_COLUMN) or "").strip():
            label = (rec.get(COL_ROW) or "").strip() or f"ردیف اکسل شماره {i + 1}"
            bad_rows.append(label)

    if bad_rows:
        shown = "، ".join(bad_rows[:15])
        more = f" (و {len(bad_rows) - 15} مورد دیگر)" if len(bad_rows) > 15 else ""
        raise ValidationError(
            f"ستون «{REQUIRED_COLUMN}» در {len(bad_rows)} ردیف خالی است: {shown}{more}. "
            "این ستون الزامی است؛ لطفاً تکمیل و دوباره آپلود کنید."
        )


def output_header_order(input_headers):
    """ترتیب ستون‌های خروجی: همان ترتیب فایل ورودی + ستون‌های خروجی که از قبل نبودند
    (هم برای اجرای اول و هم برای آپلود دوباره‌ی فایلی که قبلاً پردازش شده)."""
    order = list(input_headers)
    for col in OUTPUT_COLUMNS:
        if col not in order:
            order.append(col)
    return order


def write_result_xlsx(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "نتیجه"
    ws.append(headers)
    for rec in rows:
        ws.append([rec.get(h, "") for h in headers])
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(60, max(12, length + 2))
    wb.save(path)
